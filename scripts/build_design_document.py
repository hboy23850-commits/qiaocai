from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "设计文档-2026年.md"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT = OUTPUT_DIR / "巧裁-软件创意设计文档.docx"


def set_east_asia_font(run, font: str, size: float | None = None, bold: bool | None = None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else part
        run = paragraph.add_run(value)
        set_east_asia_font(run, "宋体", 11, is_bold)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.3)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, color in [
        ("Title", 22, "164F32"),
        ("Heading 1", 16, "164F32"),
        ("Heading 2", 14, "1B6841"),
        ("Heading 3", 12, "245F43"),
    ]:
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "巧裁 - 软件创意设计文档"
    doc.core_properties.subject = "2026年华北五省计算机应用大赛参赛作品"
    doc.core_properties.keywords = "微信小程序, 智能排料, 手工制作, 余料复用"

    separator_count = 0
    in_cover = True
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line == "---":
            separator_count += 1
            if separator_count == 1:
                doc.add_page_break()
            in_cover = False
            continue
        shot = re.fullmatch(r"\[SCREENSHOT:(.+?)\|(.+?)\]", line)
        if shot:
            image_path = ROOT / shot.group(1)
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                p.add_run().add_picture(str(image_path), height=Inches(4.8))
                caption = doc.add_paragraph(shot.group(2))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.first_line_indent = Cm(0)
                for run in caption.runs:
                    set_east_asia_font(run, "宋体", 9)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
            continue
        if line.startswith("## "):
            title = line[3:]
            p = doc.add_heading(title, level=2 if in_cover else 1)
            if in_cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(70)
            p.paragraph_format.space_after = Pt(24)
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, re.sub(r"^\d+\.\s*", "", line))
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[2:])
            continue
        p = doc.add_paragraph()
        if in_cover:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(10)
        add_inline(p, line)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
